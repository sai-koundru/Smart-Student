import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(11, 15, 25)       # Deep slate midnight
    CARD_BG = RGBColor(22, 30, 49)       # Rich dark card background
    CARD_BORDER = RGBColor(45, 55, 78)   # Subtle card border
    PRIMARY = RGBColor(139, 92, 246)     # Vibrant Violet (#8B5CF6)
    SECONDARY = RGBColor(79, 70, 229)    # Royal Indigo (#4F46E5)
    CYAN = RGBColor(6, 182, 212)         # Electric Cyan (#06B6D4)
    EMERALD = RGBColor(16, 185, 129)     # Bright Emerald (#10B981)
    AMBER = RGBColor(245, 158, 11)       # Warm Amber (#F59E0B)
    TEXT_LIGHT = RGBColor(243, 244, 246) # Crisp off-white
    TEXT_MUTED = RGBColor(156, 163, 175) # Clean gray

    def set_slide_bg(slide):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_DARK
        bg_shape.line.fill.background() # No line

        # Subtle decorative top accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PRIMARY
        bar.line.fill.background()

    def add_header(slide, title, category="GNITC PORTAL"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = CYAN
        p_c.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT
        p_t.font.name = "Arial"

    def add_card(slide, left, top, width, height, title="", subtitle="", accent_color=PRIMARY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Left accent pill
        accent = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.15), Inches(top + 0.18), Inches(0.06), Inches(height - 0.36))
        accent.fill.solid()
        accent.fill.fore_color.rgb = accent_color
        accent.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(left + 0.35), Inches(top + 0.15), Inches(width - 0.5), Inches(height - 0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(17)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT
            p.font.name = "Arial"

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(11)
            p2.font.color.rgb = accent_color
            p2.font.name = "Arial"
            p2.space_before = Pt(3)

        return tf

    # ==========================================
    # SLIDE 1: Title Slide (Hero)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Glowing decorative card in center
    hero_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = CARD_BG
    hero_card.line.color.rgb = PRIMARY
    hero_card.line.width = Pt(2)

    # University Badge
    tb = s1.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.5), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "GURU NANAK INSTITUTIONS TECHNICAL CAMPUS (AUTONOMOUS) • IBRAHIMPATNAM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.font.name = "Arial"

    # Main Project Title
    tb_title = s1.shapes.add_textbox(Inches(1.8), Inches(2.2), Inches(9.5), Inches(1.4))
    p_main = tb_title.text_frame.paragraphs[0]
    p_main.text = "GNITC PORTAL"
    p_main.font.size = Pt(48)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_LIGHT
    p_main.font.name = "Arial"

    p_sub = tb_title.text_frame.add_paragraph()
    p_sub.text = "Smart Student — Next-Generation Integrated Academic Management PWA"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = PRIMARY
    p_sub.font.name = "Arial"
    p_sub.space_before = Pt(8)

    # Description
    tb_desc = s1.shapes.add_textbox(Inches(1.8), Inches(3.9), Inches(9.5), Inches(0.8))
    p_d = tb_desc.text_frame.paragraphs[0]
    p_d.text = "A single unified cloud platform replacing fragmented systems across attendance, smart timetable, leave approvals, GNITC bus tracking, competitive exams guidance, and campus event registrations."
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = TEXT_MUTED
    p_d.font.name = "Arial"

    # Live Badge Pill
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(4.9), Inches(9.7), Inches(0.9))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(15, 23, 42)
    badge.line.color.rgb = EMERALD
    badge.line.width = Pt(1.5)

    tb_b = s1.shapes.add_textbox(Inches(2.0), Inches(5.0), Inches(9.3), Inches(0.7))
    tf_b = tb_b.text_frame
    p_b1 = tf_b.paragraphs[0]
    p_b1.text = "● LIVE CLOUD DEPLOYMENT: https://gnitc-portal.onrender.com"
    p_b1.font.size = Pt(13)
    p_b1.font.bold = True
    p_b1.font.color.rgb = EMERALD
    p_b2 = tf_b.add_paragraph()
    p_b2.text = "Mobile Progressive Web App (PWA) • Dual Dashboards (Faculty & Student) • 10 Integrated Modules"
    p_b2.font.size = Pt(11)
    p_b2.font.color.rgb = TEXT_MUTED
    p_b2.space_before = Pt(4)


    # ==========================================
    # SLIDE 2: Problem vs Solution
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "The Challenge vs The GNITC Portal Solution", "Project Motivation")

    # Left Card: The Problem
    tf2_l = add_card(s2, 0.8, 1.8, 5.6, 5.0, "The Fragmented Campus Problem", "Legacy & Multi-App Inefficiencies", AMBER)
    items_l = [
        ("Disjointed Platforms", "Students toggle 4-5 different platforms for attendance, notes, assignments, and announcements."),
        ("Manual Leave Friction", "Leave letters take days to get signed by HODs and verified with parents over phone calls."),
        ("Bus Route Uncertainty", "Students commuting 50-60 km across Hyderabad lack exact bus stop arrival & departure times."),
        ("Missed Opportunities", "Crucial hackathons, internship alerts, and exam deadlines get lost in WhatsApp group clutter.")
    ]
    for title, desc in items_l:
        p = tf2_l.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right Card: The Solution
    tf2_r = add_card(s2, 6.9, 1.8, 5.6, 5.0, "The Integrated GNITC Ecosystem", "Single Unified Mobile Solution", EMERALD)
    items_r = [
        ("All-in-One Dashboard", "Combines timetable, marks, QR attendance, assignments, and alerts into one sleek portal."),
        ("Paperless Digital Leave", "Instant 2-tier approval where parents approve via a secure web link, followed by teacher sign-off."),
        ("Live Route Guide", "Dedicated 8-route GNITC bus guide detailing timings from Dilsukhnagar, Secunderabad, Suchitra, etc."),
        ("Career Roadmaps & Events", "Comprehensive pre/post-engineering competitive exam guides and categorized campus event boards.")
    ]
    for title, desc in items_r:
        p = tf2_r.add_paragraph()
        p.text = f"✔ {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED


    # ==========================================
    # SLIDE 3: System Architecture & Stack
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Modern Full-Stack Architecture & Technology", "Engineering Specifications")

    col_w = 3.65
    gap = 0.35

    # Column 1: Frontend
    tf3_1 = add_card(s3, 0.8, 1.8, col_w, 5.0, "Frontend Layer", "React 18 + Vite + Tailwind CSS", PRIMARY)
    p = tf3_1.add_paragraph()
    p.text = "• React 18 SPA with React Router v6\n• Tailwind CSS modern glassmorphic theme\n• Lucide React icons library\n• HTML5-QRCode scanner module\n• QRCode.react SVG session generator\n• Mobile-first responsive touch layout\n• Stored JWT auth with reactive context"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)

    # Column 2: Backend
    tf3_2 = add_card(s3, 0.8 + col_w + gap, 1.8, col_w, 5.0, "Backend & API Layer", "Node.js + Express + SQLite", CYAN)
    p = tf3_2.add_paragraph()
    p.text = "• Express.js REST API with clean modular routes\n• SQLite with Better-SQLite3 native driver\n• BcryptJS salted password hashing\n• JWT Token authentication with 7-day expiry\n• Role-based authorization middleware\n• Static production asset serving from root\n• Automated seed generator & DB recovery"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)

    # Column 3: Cloud & PWA
    tf3_3 = add_card(s3, 0.8 + 2 * (col_w + gap), 1.8, col_w, 5.0, "Deployment & PWA", "Render.com + Service Worker", EMERALD)
    p = tf3_3.add_paragraph()
    p.text = "• Web Manifest (standalone full-screen app)\n• Background service worker caching (v2)\n• Custom GNITC vector launcher icons\n• Zero-app-store 1-tap browser installation\n• Continuous deployment from GitHub repo\n• 24/7 Global accessibility via HTTPS\n• High performance <3s production Vite build"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)


    # ==========================================
    # SLIDE 4: Academic Core Modules
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Smart Academic Core: Daily Student Workflows", "Core Functionalities")

    card_w = 5.6
    card_h = 2.35
    row_gap = 0.3

    # Card 1: Attendance
    tf4_1 = add_card(s4, 0.8, 1.8, card_w, card_h, "QR-Based Smart Attendance", "Fast, Anti-Proxy Classroom Tracking", CYAN)
    p = tf4_1.add_paragraph()
    p.text = "Faculty generates dynamic, time-limited QR codes on classroom screens. Students scan using camera to mark attendance instantly. Real-time percentage tracking with color thresholds (>75%)."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)

    # Card 2: Timetable
    tf4_2 = add_card(s4, 6.9, 1.8, card_w, card_h, "Dynamic Timetable System", "Live Schedule & Classroom Allocation", PRIMARY)
    p = tf4_2.add_paragraph()
    p.text = "Weekly day-by-day timetable grid highlighting current day's classes, subject codes, lab rooms, and assigned faculty. Faculty dashboard provides instant slot creation and rescheduling."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)

    # Card 3: Assignments
    tf4_3 = add_card(s4, 0.8, 1.8 + card_h + row_gap, card_w, card_h, "Assignment Management Hub", "Digital Homework & Feedback", AMBER)
    p = tf4_3.add_paragraph()
    p.text = "Teachers upload homework with deadlines and maximum marks. Students submit text/solutions directly in the portal. Real-time status badges (Pending, Submitted, Graded) with remarks."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)

    # Card 4: Marks
    tf4_4 = add_card(s4, 6.9, 1.8 + card_h + row_gap, card_w, card_h, "Marks & Academic Performance", "Multi-Exam Visual Analytics", EMERALD)
    p = tf4_4.add_paragraph()
    p.text = "Detailed scorecards across Midterms, Quizzes, Internals, and Semester Finals. Visual bar charts give students clear insights into course strengths, percentiles, and subject performance."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)


    # ==========================================
    # SLIDE 5: GNITC Bus Track System
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "College Bus Track: Built for GNITC Ibrahimpatnam", "Campus Transport Feature")

    # Left Details Card
    tf5_l = add_card(s5, 0.8, 1.8, 4.4, 5.0, "GNITC Fleet Highlights", "Covering 50-60 KM Radius of Hyderabad", CYAN)
    p = tf5_l.add_paragraph()
    p.text = "• Location: Ibrahimpatnam, R.R. Dist, Telangana\n• Fleet Size: 50+ AC & Non-AC Buses\n• Transport Helpline: +91 84990 71144\n• Morning Arrival: 8:45 AM – 9:00 AM\n• Evening Departure: 4:30 PM Sharp\n• Mandatory Policy: Valid bus pass & ID card\n• Real-Time Search: Students can filter routes instantly by locality name."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)

    # Right Routes Grid Card
    tf5_r = add_card(s5, 5.6, 1.8, 6.9, 5.0, "Documented Point-to-Point (P2P) Routes", "Exact Morning Pickup & Evening Timings", PRIMARY)
    routes = [
        ("Bus 1: Dilsukhnagar", "7:30 AM", "Dilsukhnagar → Kothapet → LB Nagar → Hayathnagar → GNITC (8:50 AM)"),
        ("Bus 2: Secunderabad", "7:15 AM", "Secunderabad → Tarnaka → Uppal → Boduppal → GNITC (8:50 AM)"),
        ("Bus 3: Mehdipatnam", "7:00 AM", "Mehdipatnam → Attapur → Rajendranagar → Shamshabad → GNITC (8:45 AM)"),
        ("Bus 4: Suchitra", "6:45 AM", "Suchitra → Alwal → Malkajgiri → ECIL → Nacharam → GNITC (8:50 AM)"),
        ("Bus 5: Naredmet", "7:00 AM", "Naredmet → Sainikpuri → ECIL → AS Rao Nagar → GNITC (8:50 AM)"),
        ("Bus 6: Kukatpally", "6:45 AM", "Kukatpally → JNTU → KPHB → Miyapur → Chandanagar → GNITC (8:50 AM)"),
        ("Bus 7: Malakpet", "7:15 AM", "Malakpet → Chaderghat → Vanasthalipuram → Hayathnagar → GNITC (8:45 AM)"),
        ("Bus 8: Medipally", "7:15 AM", "Medipally → Uppal Depot → Pirzadiguda → Nagole → GNITC (8:50 AM)")
    ]
    for r_title, r_time, r_stops in routes:
        p = tf5_r.add_paragraph()
        p.text = f"🚌 {r_title} ({r_time}): "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = r_stops
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED


    # ==========================================
    # SLIDE 6: Digital Leave Workflow
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Digital Leave Application: Dual-Approval Architecture", "Security & Verification")

    step_w = 2.7
    step_gap = 0.25

    steps = [
        ("Step 1: Student Applies", "Student fills leave dates and explanation of absence directly in the portal. A unique UUID verification token is generated.", PRIMARY),
        ("Step 2: Parent Link Generated", "The portal produces a dedicated standalone URL (e.g. /parent-approval/uuid). Student shares it with their parent via WhatsApp/SMS.", CYAN),
        ("Step 3: 1-Tap Parent Decision", "Parents open the secure page on their phone without logging in. They view absence reasons and tap 'Approve' or 'Reject' with comments.", EMERALD),
        ("Step 4: Faculty Final Review", "Class Teacher views the live status in 'Leave Approvals'. Leave is officially granted only when BOTH parent and teacher grant approval.", AMBER)
    ]

    for i, (stitle, sdesc, scolor) in enumerate(steps):
        tf_s = add_card(s6, 0.8 + i * (step_w + step_gap), 1.8, step_w, 3.4, f"0{i+1}", stitle, scolor)
        p = tf_s.add_paragraph()
        p.text = sdesc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # Bottom Summary Callout
    tf6_b = add_card(s6, 0.8, 5.5, 11.7, 1.4, "Why This Architecture Matters for GNITC", "Zero Proxy Leaves & Complete Accountability", EMERALD)
    p = tf6_b.add_paragraph()
    p.text = "✔ Eliminates forged paper letters  •  ✔ Saves faculty hours spent on phone verification  •  ✔ Gives parents immediate visibility  •  ✔ Complete timestamped audit trail in SQLite"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(6)


    # ==========================================
    # SLIDE 7: Career & Competitive Exams Guide
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Comprehensive Exams Guide: 25+ Career Roadmaps", "Student Career Guidance")

    # Left: Before B.Tech
    tf7_l = add_card(s7, 0.8, 1.8, 5.6, 5.0, "1. Pre-Engineering Pathways", "College Admissions & Central Govt Exams", CYAN)
    exams_pre = [
        ("JEE Main & Advanced", "IITs, NITs, IIITs — Elite Software & Core Engineering"),
        ("TS & AP EAMCET", "State Engineering Colleges (Telangana & Andhra Pradesh)"),
        ("BITSAT & COMEDK", "BITS Pilani & Top Private Karnataka Engineering Colleges"),
        ("SAT & ACT", "Undergraduate Admissions in Top USA & Canadian Universities"),
        ("IELTS & TOEFL", "English Proficiency Standard for Overseas University Studies"),
        ("SSC CHSL & RRB NTPC", "Central Govt Clerks, Postal Assistants, Railway Officers")
    ]
    for e_name, e_desc in exams_pre:
        p = tf7_l.add_paragraph()
        p.text = f"★ {e_name}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(9)
        run = p.add_run()
        run.text = e_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right: After B.Tech
    tf7_r = add_card(s7, 6.9, 1.8, 5.6, 5.0, "2. Post-Engineering Pathways", "Higher Studies, PSU & Elite Govt Roles", PRIMARY)
    exams_post = [
        ("GATE Exam", "M.Tech at IITs/IISc and Direct PSU Recruitment (ONGC, BHEL)"),
        ("CAT & GMAT", "MBA at IIMs, Harvard, Stanford — Executive Leadership"),
        ("GRE & TOEFL/IELTS", "Master of Science (MS) / PhD in USA, Germany, Europe"),
        ("UPSC Civil Services (CSE)", "IAS, IPS, IFS, IRS — Nation's Premier Bureaucracy"),
        ("SSC CGL & Bank PO (SBI/IBPS)", "Income Tax Inspector, Central Auditor, Bank Officer"),
        ("ISRO Scientist & DRDO SET", "Space Scientist & Defence Research Scientist Positions")
    ]
    for e_name, e_desc in exams_post:
        p = tf7_r.add_paragraph()
        p.text = f"★ {e_name}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(9)
        run = p.add_run()
        run.text = e_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED


    # ==========================================
    # SLIDE 8: Event Management & Campus Life
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Campus Event Management & Student Opportunities", "Student Development")

    col8_w = 3.65
    col8_gap = 0.35

    # Card 1: Hackathons & Tech
    tf8_1 = add_card(s8, 0.8, 1.8, col8_w, 5.0, "Tech & Hackathons", "Competitions & Summer Schools", PRIMARY)
    p = tf8_1.add_paragraph()
    p.text = "• GNITC Code Sprint 2026\n  Annual flagship 24h coding marathon\n  Registration: ₹100 • Cash awards\n\n• TCS CodeVita Global\n  National competitive coding challenge\n  Free entry • Fast-track interview\n\n• Amazon ML Summer School\n  Deep-dive machine learning training\n  Free selection for 3rd/4th year students"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)

    # Card 2: Workshops & Seminars
    tf8_2 = add_card(s8, 0.8 + col8_w + col8_gap, 1.8, col8_w, 5.0, "Skill Workshops", "Hands-On Practical Training", CYAN)
    p = tf8_2.add_paragraph()
    p.text = "• AI & LLM Practical Workshop\n  Building modern AI agents & apps\n  Registration: ₹200 • Certificate\n\n• Campus Recruitment Drive\n  Resume screening & mock interviews\n  Free entry • Placement preparation\n\n• Cloud DevOps Bootcamp\n  Docker, Kubernetes & CI/CD pipelines\n  Conducted with industry experts"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)

    # Card 3: Cultural & Social
    tf8_3 = add_card(s8, 0.8 + 2 * (col8_w + col8_gap), 1.8, col8_w, 5.0, "Cultural & Fests", "Extra-Curricular Activities", AMBER)
    p = tf8_3.add_paragraph()
    p.text = "• Sargam 2026 Cultural Fest\n  GNITC Inter-college cultural carnival\n  Music, Dance, Drama, Fine Arts\n  Entry: ₹50 • Trophy prizes\n\n• Annual Sports Meet\n  Cricket, Football, Basketball\n  Inter-department tournament\n\n• Faculty Control Panel\n  Teachers can publish, edit deadlines,\n  and post application links in 10 seconds"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)


    # ==========================================
    # SLIDE 9: Live Demo, Summary & Q&A
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Live Access, Credentials & Project Summary", "Deployment & Verification")

    # Left Card: Live Deployment & Test Logins
    tf9_l = add_card(s9, 0.8, 1.8, 5.6, 5.0, "Live Platform Access", "Deployed On Render Cloud (24/7 Live)", EMERALD)
    p = tf9_l.add_paragraph()
    p.text = "🔗 Live URL: https://gnitc-portal.onrender.com\n\n" \
             "👤 STUDENT TEST CREDENTIALS:\n" \
             "• Username: sanjana.sambu@test.com\n" \
             "• Password:  0987654321\n" \
             "• Access: Timetable, Bus Track, QR Scanner,\n" \
             "  Leave Application, Exams Guide, Events\n\n" \
             "👨‍🏫 FACULTY TEST CREDENTIALS:\n" \
             "• Username: sai_koundru@test.com\n" \
             "• Password:  1234567890\n" \
             "• Access: QR Generator, Leave Approvals,\n" \
             "  Marks Entry, Event Publishing, Alerts"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(12)

    # Right Card: Project Impact & Future Scope
    tf9_r = add_card(s9, 6.9, 1.8, 5.6, 5.0, "Project Impact & Future Scope", "Transforming Campus Administration", PRIMARY)
    impacts = [
        ("100% Paperless Operations", "Digitizes attendance, leaves, and assignment cycles across departments."),
        ("Zero Install Barrier", "Progressive Web App installs instantly on Android & iOS without app store delays."),
        ("Enhanced Parent-College Trust", "Direct 1-tap parental leave approval links remove guesswork and proxies."),
        ("Future Enhancements", "GPS integration for live GNITC bus location tracking, fee payment gateway, and library book management.")
    ]
    for title, desc in impacts:
        p = tf9_r.add_paragraph()
        p.text = f"🚀 {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Save to project folder
    output_path = r"C:\Users\AWS-SAI\.gemini\antigravity\scratch\smart-student\GNITC_Portal_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_presentation()
